#!/usr/bin/env python
"""Extract tile metadata and clean tile images from a catalogue PDF."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import logging
import math
import re
import shutil
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

import pdfplumber
from pdfminer.pdftypes import resolve1
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from PIL import Image, ImageStat


DIM_RE = re.compile(r"(?P<a>\d{3,4})\s*[xX×*]\s*(?P<b>\d{3,4})\s*(?:m\s*m|mm|MM)?")
SAFE_RE = re.compile(r"[^A-Za-z0-9._-]+")
HEADER_WORDS = {
    "SIZE",
    "PGVT",
    "VITRIFIED",
    "TILES",
    "DIGITAL",
    "WALL",
    "FLOOR",
    "STEP",
    "RISER",
    "RESER",
    "SERIES",
    "GLOSSY",
    "MATT",
    "FINISH",
    "STRIP",
}


@dataclass
class DimensionHint:
    length: int
    width: int
    x: float
    top: float
    text: str


@dataclass
class TileRecord:
    title: str
    length_mm: int | None
    width_mm: int | None
    image_name: str
    page: int
    image_object: str
    source_width_px: int
    source_height_px: int
    output_width_px: int
    output_height_px: int
    aspect_delta_pct: float | None
    extraction_method: str
    warnings: list[str]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract tile images and metadata from a PDF catalogue.")
    parser.add_argument("input_pdf", type=Path, help="Path to the tile catalogue PDF")
    parser.add_argument("--output-dir", type=Path, default=Path("submission"), help="Output folder")
    parser.add_argument("--no-upscale", action="store_true", help="Keep native embedded image sizes")
    parser.add_argument("--log-level", default="INFO", choices=["DEBUG", "INFO", "WARNING", "ERROR"])
    return parser.parse_args()


def slugify(value: str, max_len: int = 70) -> str:
    value = value.strip().replace("&", "and")
    value = SAFE_RE.sub("_", value)
    value = re.sub(r"_+", "_", value).strip("._-")
    return (value or "tile")[:max_len]


def word_center(word: dict[str, Any]) -> tuple[float, float]:
    return ((word["x0"] + word["x1"]) / 2, (word["top"] + word["bottom"]) / 2)


def image_display_aspect(image_obj: dict[str, Any]) -> float:
    width = float(image_obj["width"])
    height = float(image_obj["height"])
    return width / height if height else 0.0


def physical_ratio(length: int, width: int) -> float:
    small = max(1, min(length, width))
    large = max(length, width)
    return large / small


def image_ratio(image_obj: dict[str, Any]) -> float:
    ratio = image_display_aspect(image_obj)
    if ratio <= 0:
        src_w, src_h = image_obj.get("srcsize", (1, 1))
        ratio = src_w / max(1, src_h)
    return max(ratio, 1 / ratio)


def parse_dimension_hints(words: list[dict[str, Any]]) -> list[DimensionHint]:
    hints: list[DimensionHint] = []
    seen: set[tuple[int, int, int, int]] = set()
    for word in words:
        text = word["text"].replace(" ", "")
        match = DIM_RE.search(text)
        if not match:
            continue
        length = int(match.group("a"))
        width = int(match.group("b"))
        key = (length, width, round(word["x0"]), round(word["top"]))
        if key in seen:
            continue
        seen.add(key)
        hints.append(DimensionHint(length, width, word["x0"], word["top"], word["text"]))
    return hints


def first_label_line(image_obj: dict[str, Any], words: list[dict[str, Any]]) -> str | None:
    bottom = image_obj["bottom"]
    x0 = image_obj["x0"] - 6
    x1 = image_obj["x1"] + 6
    candidates = [
        w
        for w in words
        if bottom - 2 <= w["top"] <= bottom + 46
        and w["x1"] >= x0
        and w["x0"] <= x1
    ]
    if not candidates:
        return None

    candidates.sort(key=lambda w: (w["top"], w["x0"]))
    lines: list[list[dict[str, Any]]] = []
    for word in candidates:
        placed = False
        for line in lines:
            if abs(line[0]["top"] - word["top"]) <= 4:
                line.append(word)
                placed = True
                break
        if not placed:
            lines.append([word])

    for line in lines:
        line.sort(key=lambda w: w["x0"])
        text = " ".join(w["text"] for w in line)
        text = re.sub(r"\s+", " ", text).strip()
        if is_valid_title(text, image_obj):
            return text
    return None


def is_valid_title(title: str, image_obj: dict[str, Any]) -> bool:
    cleaned = title.strip(" .:|")
    if not cleaned:
        return False
    upper = cleaned.upper()
    if upper in HEADER_WORDS:
        return False
    if DIM_RE.search(cleaned.replace(" ", "")):
        return False
    if re.fullmatch(r"\d{1,2}", cleaned):
        return False
    src_w, src_h = image_obj.get("srcsize", (0, 0))
    display_area = image_obj["width"] * image_obj["height"]
    if re.fullmatch(r"\d{1,3}", cleaned) and (display_area > 90000 or src_w > 800 or src_h > 800):
        return False
    return bool(re.search(r"[A-Za-z0-9]", cleaned))


def is_candidate_image(image_obj: dict[str, Any]) -> bool:
    src_w, src_h = image_obj.get("srcsize", (0, 0))
    display_area = image_obj["width"] * image_obj["height"]
    if src_w < 120 or src_h < 120:
        return False
    if display_area < 4500:
        return False
    ratio = image_ratio(image_obj)
    if not 0.9 <= ratio <= 7.2:
        return False
    return True


def choose_dimension(image_obj: dict[str, Any], hints: list[DimensionHint]) -> tuple[DimensionHint | None, float | None]:
    if not hints:
        return None, None
    ratio = image_ratio(image_obj)
    image_cx = (image_obj["x0"] + image_obj["x1"]) / 2

    def score(hint: DimensionHint) -> tuple[float, float]:
        ratio_score = abs(math.log(ratio / physical_ratio(hint.length, hint.width)))
        same_half_bonus = 0 if (image_cx < 595 and hint.x < 595) or (image_cx >= 595 and hint.x >= 595) else 0.08
        return ratio_score + same_half_bonus, abs(image_cx - hint.x)

    best = min(hints, key=score)
    delta_pct = abs(ratio - physical_ratio(best.length, best.width)) / physical_ratio(best.length, best.width) * 100
    return best, delta_pct


def reconcile_dimension_with_tile_aspect(
    title: str,
    image_obj: dict[str, Any],
    dimension: DimensionHint | None,
) -> tuple[int | None, int | None, float | None, list[str]]:
    if dimension is None:
        return None, None, None, ["missing dimensions"]

    length, width = dimension.length, dimension.width
    ratio = image_ratio(image_obj)
    delta_pct = abs(ratio - physical_ratio(length, width)) / physical_ratio(length, width) * 100
    warnings: list[str] = []

    # Wall-tile sets often include a square floor/companion tile in the same product group.
    # The catalogue header states the wall size, while the embedded tile image is square.
    # Use the stated smaller side to derive the square companion size and flag it in the log.
    smaller_side = min(length, width)
    larger_side = max(length, width)
    if delta_pct > 8 and abs(ratio - 1.0) <= 0.08 and smaller_side == 300 and larger_side in {450, 600}:
        length = width = 300
        delta_pct = abs(ratio - 1.0) * 100
        warnings.append("derived square companion size from catalogue section and image aspect")

    if delta_pct > 8:
        warnings.append(f"aspect differs from dimensions by {delta_pct:.1f}%")
    return length, width, delta_pct, warnings


def decode_embedded_image(image_obj: dict[str, Any]) -> tuple[Image.Image, str]:
    data = image_obj["stream"].get_data()
    try:
        image = Image.open(io.BytesIO(data))
        image.load()
        return image.convert("RGB"), "embedded"
    except Exception:
        attrs = image_obj["stream"].attrs
        width = int(attrs.get("Width", image_obj["srcsize"][0]))
        height = int(attrs.get("Height", image_obj["srcsize"][1]))
        colorspace = resolve1(attrs.get("ColorSpace"))
        bits = int(attrs.get("BitsPerComponent", 8))

        if bits == 8 and isinstance(colorspace, list) and str(colorspace[0]).endswith("Indexed'"):
            palette = resolve1(colorspace[3])
            image = Image.frombytes("P", (width, height), data)
            padded_palette = bytes(palette) + (b"\x00" * max(0, 768 - len(palette)))
            image.putpalette(padded_palette[:768])
            return image.convert("RGB"), "embedded-indexed"

        if bits == 8 and len(data) == width * height:
            image = Image.frombytes("L", (width, height), data)
            return image.convert("RGB"), "embedded-gray"

        if bits == 8 and len(data) == width * height * 3:
            raw = Image.frombytes("RGB", (width, height), data)
            return raw.convert("RGB"), "embedded-raw"

        raise ValueError(f"unsupported raw image stream: {width}x{height}, bits={bits}, bytes={len(data)}")


def maybe_upscale(image: Image.Image, length: int | None, width: int | None, no_upscale: bool) -> Image.Image:
    if no_upscale or not length or not width:
        return image
    if image.width >= image.height:
        required_w, required_h = max(length, width), min(length, width)
    else:
        required_w, required_h = min(length, width), max(length, width)
    scale = max(1.0, required_w / image.width, required_h / image.height)
    if scale <= 1.0001:
        return image
    new_size = (math.ceil(image.width * scale), math.ceil(image.height * scale))
    return image.resize(new_size, Image.Resampling.LANCZOS)


def image_blank_score(image: Image.Image) -> float:
    stat = ImageStat.Stat(image.convert("L"))
    return float(stat.stddev[0])


def write_excel(records: list[TileRecord], output_xlsx: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tiles"
    headers = ["Title", "Length (mm)", "Width (mm)", "Image Name"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for record in records:
        ws.append([record.title, record.length_mm, record.width_mm, record.image_name])
    widths = [42, 14, 14, 48]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    wb.save(output_xlsx)


def write_csv_log(records: list[TileRecord], output_csv: Path) -> None:
    fieldnames = list(asdict(records[0]).keys()) if records else ["title"]
    with output_csv.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for record in records:
            row = asdict(record)
            row["warnings"] = "; ".join(record.warnings)
            writer.writerow(row)


def validate(records: list[TileRecord], image_dir: Path, output_xlsx: Path) -> dict[str, Any]:
    warnings: list[str] = []
    informational: list[str] = []
    image_files = sorted(p.name for p in image_dir.glob("*.png"))
    image_names = [r.image_name for r in records]

    missing_metadata = [r.image_name for r in records if not r.title or not r.length_mm or not r.width_mm]
    missing_images = [name for name in image_names if not (image_dir / name).exists()]
    duplicate_records = [
        "|".join(map(str, key))
        for key, count in Counter((r.title, r.length_mm, r.width_mm) for r in records).items()
        if count > 1
    ]
    duplicate_image_names = [name for name, count in Counter(image_names).items() if count > 1]

    bad_aspect: list[str] = []
    aspect_notices: list[str] = []
    small_images: list[str] = []
    blank_images: list[str] = []
    below_required: list[str] = []

    for record in records:
        path = image_dir / record.image_name
        if not path.exists():
            continue
        with Image.open(path) as img:
            if min(img.size) < 80:
                small_images.append(record.image_name)
            low_variation_ok = re.search(r"WHITE|PLAIN|IVORY|BLACK|MATT|SUPER", record.title, re.IGNORECASE)
            if image_blank_score(img) < 0.05 and not low_variation_ok:
                blank_images.append(record.image_name)
            if record.length_mm and record.width_mm:
                img_ratio = max(img.width / img.height, img.height / img.width)
                target_ratio = physical_ratio(record.length_mm, record.width_mm)
                delta = abs(img_ratio - target_ratio) / target_ratio * 100
                if delta > 15:
                    bad_aspect.append(f"{record.image_name}: {delta:.1f}%")
                elif delta > 8:
                    aspect_notices.append(f"{record.image_name}: {delta:.1f}%")
                if img.width >= img.height:
                    req_w, req_h = max(record.length_mm, record.width_mm), min(record.length_mm, record.width_mm)
                else:
                    req_w, req_h = min(record.length_mm, record.width_mm), max(record.length_mm, record.width_mm)
                if img.width < req_w or img.height < req_h:
                    below_required.append(record.image_name)

    workbook_rows = None
    workbook_headers = []
    if output_xlsx.exists():
        wb = load_workbook(output_xlsx, read_only=True, data_only=True)
        ws = wb.active
        workbook_headers = [ws.cell(1, col).value for col in range(1, 5)]
        workbook_rows = max(0, ws.max_row - 1)
        wb.close()

    if missing_metadata:
        warnings.append(f"Missing metadata for {len(missing_metadata)} records")
    if duplicate_records:
        informational.append(f"Duplicate title/dimension combinations present in catalogue: {len(duplicate_records)}")
    if missing_images:
        warnings.append(f"Missing image files: {len(missing_images)}")
    if duplicate_image_names:
        warnings.append(f"Duplicate image names: {len(duplicate_image_names)}")
    if bad_aspect:
        warnings.append(f"Severe aspect-ratio outliers: {len(bad_aspect)}")
    if aspect_notices:
        informational.append(f"Moderate aspect-ratio notices: {len(aspect_notices)}")
    if below_required:
        warnings.append(f"Images below physical-pixel target: {len(below_required)}")
    if small_images:
        warnings.append(f"Suspiciously small images: {len(small_images)}")
    if blank_images:
        warnings.append(f"Suspiciously blank images: {len(blank_images)}")
    if workbook_rows != len(records):
        warnings.append("Excel row count does not match record count")
    if len(image_files) != len(records):
        warnings.append("Image count does not match record count")
    if workbook_headers != ["Title", "Length (mm)", "Width (mm)", "Image Name"]:
        warnings.append("Excel headers do not match required columns")

    return {
        "tile_records": len(records),
        "image_files": len(image_files),
        "excel_rows": workbook_rows,
        "required_headers": ["Title", "Length (mm)", "Width (mm)", "Image Name"],
        "excel_headers": workbook_headers,
        "missing_metadata": missing_metadata,
        "duplicate_title_dimension_records": duplicate_records,
        "duplicate_image_names": duplicate_image_names,
        "missing_images": missing_images,
        "aspect_ratio_outliers": bad_aspect,
        "aspect_ratio_notices": aspect_notices,
        "below_required_pixel_target": below_required,
        "suspiciously_small_images": small_images,
        "suspiciously_blank_images": blank_images,
        "warnings": warnings,
        "informational": informational,
        "passed": not warnings,
    }


def stable_image_name(page_num: int, seq: int, title: str, image: Image.Image) -> str:
    digest = hashlib.sha1(f"{page_num}-{seq}-{title}-{image.width}x{image.height}".encode("utf-8")).hexdigest()[:8]
    return f"p{page_num:03d}_{seq:03d}_{slugify(title)}_{digest}.png"


def extract_tiles(input_pdf: Path, output_dir: Path, no_upscale: bool) -> list[TileRecord]:
    image_dir = output_dir / "images"
    image_dir.mkdir(parents=True, exist_ok=True)
    records: list[TileRecord] = []

    with pdfplumber.open(str(input_pdf)) as pdf:
        for page in pdf.pages:
            words = page.extract_words(x_tolerance=2, y_tolerance=2, keep_blank_chars=False, use_text_flow=False)
            hints = parse_dimension_hints(words)
            page_seq = 0
            for image_obj in page.images:
                if not is_candidate_image(image_obj):
                    continue
                title = first_label_line(image_obj, words)
                if not title:
                    continue
                dimension, aspect_delta = choose_dimension(image_obj, hints)
                length, width, aspect_delta, warnings = reconcile_dimension_with_tile_aspect(title, image_obj, dimension)

                try:
                    source_img, method = decode_embedded_image(image_obj)
                except Exception as exc:
                    logging.warning("Skipping page %s image %s: cannot decode (%s)", page.page_number, image_obj.get("name"), exc)
                    continue

                final_img = maybe_upscale(source_img, length, width, no_upscale)
                page_seq += 1
                image_name = stable_image_name(page.page_number, page_seq, title, final_img)
                final_img.save(image_dir / image_name, "PNG", optimize=True)

                records.append(
                    TileRecord(
                        title=title,
                        length_mm=length,
                        width_mm=width,
                        image_name=image_name,
                        page=page.page_number,
                        image_object=str(image_obj.get("name", "")),
                        source_width_px=source_img.width,
                        source_height_px=source_img.height,
                        output_width_px=final_img.width,
                        output_height_px=final_img.height,
                        aspect_delta_pct=round(aspect_delta, 3) if aspect_delta is not None else None,
                        extraction_method=method,
                        warnings=warnings,
                    )
                )

    records.sort(key=lambda r: (r.page, r.image_object, r.title))
    return records


def prepare_output_dir(output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    for child in ["images", "output.xlsx", "extraction_log.csv", "validation_report.json"]:
        path = output_dir / child
        if path.is_dir():
            shutil.rmtree(path)
        elif path.exists():
            path.unlink()


def main() -> int:
    args = parse_args()
    logging.basicConfig(level=getattr(logging, args.log_level), format="%(levelname)s: %(message)s")
    input_pdf = args.input_pdf.resolve()
    output_dir = args.output_dir.resolve()
    if not input_pdf.exists():
        raise FileNotFoundError(f"Input PDF not found: {input_pdf}")

    prepare_output_dir(output_dir)
    records = extract_tiles(input_pdf, output_dir, args.no_upscale)
    output_xlsx = output_dir / "output.xlsx"
    write_excel(records, output_xlsx)
    write_csv_log(records, output_dir / "extraction_log.csv")
    report = validate(records, output_dir / "images", output_xlsx)
    (output_dir / "validation_report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")

    logging.info("Extracted %s tile records", len(records))
    logging.info("Images: %s", output_dir / "images")
    logging.info("Excel: %s", output_xlsx)
    if report["warnings"]:
        for warning in report["warnings"]:
            logging.warning(warning)
    else:
        logging.info("Validation passed")
    for note in report.get("informational", []):
        logging.info(note)
    return 0 if report["passed"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
